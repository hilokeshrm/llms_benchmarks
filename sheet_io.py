"""Read/write layer for dgx_llm_benchmark.xlsx.

Deliberately owns no schema of its own: headers are read from row 1 of each
result sheet, and derived formulas are copied from a row that already has them
(via openpyxl's Translator). So build_workbook.py stays the single source of
truth - add a column there and this module picks it up with no changes.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import Font

BODY = Font(name="Arial", size=10)


# --------------------------------------------------------------- reading ---
def _kv_sheet(ws, start_row=4):
    """Field/Value sheets (Run_Config, Hardware_Env) -> {field: value}."""
    out = {}
    for r in range(start_row, ws.max_row + 1):
        k = ws.cell(row=r, column=1).value
        if not k or str(k).startswith("#"):
            continue
        out[str(k).strip()] = ws.cell(row=r, column=2).value
    return out


def load_run_config(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        return _kv_sheet(wb["Run_Config"])
    finally:
        wb.close()


def load_hardware_env(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        return _kv_sheet(wb["Hardware_Env"])
    finally:
        wb.close()


def load_models(path, priority=None, only=None, require_ollama_tag=False):
    """Models sheet -> list of dicts, filtered."""
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb["Models"]
        rows = list(ws.iter_rows(min_row=4, values_only=True))
        headers = [h for h in rows[0] if h]
        out = []
        for raw in rows[1:]:
            rec = dict(zip(headers, raw))
            key = rec.get("model_key")
            if not key or not str(key).strip():
                continue
            rec["ollama_tag"] = (rec.get("ollama_tag") or "").strip()
            if only and key not in only:
                continue
            if priority and str(rec.get("priority") or "").upper() not in priority:
                continue
            if require_ollama_tag and not rec["ollama_tag"]:
                continue
            out.append(rec)
        return out
    finally:
        wb.close()


# --------------------------------------------------------------- writing ---
def _headers(ws, row=1):
    return {c.value: c.column for c in ws[row] if c.value}


def _formula_columns(ws, header_row, probe_row):
    """Columns whose probe_row cell holds a formula -> {col_idx: formula_string}."""
    return {
        c.column: c.value
        for c in ws[probe_row]
        if isinstance(c.value, str) and c.value.startswith("=")
    }


def _first_free_row(ws, key_col, header_row, example_prefix="EXAMPLE"):
    """First row with no run_id. The pre-seeded EXAMPLE row is treated as free."""
    r = header_row + 1
    while True:
        v = ws.cell(row=r, column=key_col).value
        if v is None or str(v).strip() == "" or str(v).startswith(example_prefix):
            return r
        r += 1


def append_rows(path, sheet_name, records, key_field, header_row=1):
    """Append dicts to a result sheet, preserving/extending derived formulas.

    Only keys matching a header are written; unknown keys are ignored, and
    formula columns are never overwritten by data.
    """
    records = [r for r in records if r]
    if not records:
        return 0
    wb = load_workbook(path)  # keep formulas
    try:
        ws = wb[sheet_name]
        hdr = _headers(ws, header_row)
        if key_field not in hdr:
            raise KeyError(f"{sheet_name} has no '{key_field}' column")
        probe = header_row + 1
        formulas = _formula_columns(ws, header_row, probe)
        row = _first_free_row(ws, hdr[key_field], header_row)

        for rec in records:
            # Clear a leftover EXAMPLE row's data cells before writing over it.
            if str(ws.cell(row=row, column=hdr[key_field]).value or "").startswith("EXAMPLE"):
                for col in hdr.values():
                    if col not in formulas:
                        ws.cell(row=row, column=col).value = None
            for name, col in hdr.items():
                if col in formulas:
                    continue  # derived - leave the formula alone
                if name in rec:
                    c = ws.cell(row=row, column=col)
                    c.value = rec[name]
                    c.font = BODY
            # Extend formulas if we ran past the pre-filled block.
            for col, template in formulas.items():
                cell = ws.cell(row=row, column=col)
                if cell.value is None:
                    src = f"{cell.column_letter}{probe}"
                    dst = f"{cell.column_letter}{row}"
                    cell.value = Translator(template, origin=src).translate_formula(dst)
            row += 1
        wb.save(path)
        return len(records)
    finally:
        wb.close()


def append_perf(path, records):
    return append_rows(path, "Perf_Results", records, "run_id")


def append_quality(path, records):
    return append_rows(path, "Quality_Results", records, "run_id")


def append_pull_log(path, records):
    return append_rows(path, "Ollama_Pull_Log", records, "model_key", header_row=4)


def update_hardware_env(path, values):
    """Write auto-detected environment facts into Hardware_Env (only non-empty)."""
    wb = load_workbook(path)
    try:
        ws = wb["Hardware_Env"]
        written = 0
        for r in range(4, ws.max_row + 1):
            k = ws.cell(row=r, column=1).value
            if k in values and values[k] not in (None, ""):
                ws.cell(row=r, column=2).value = values[k]
                written += 1
        wb.save(path)
        return written
    finally:
        wb.close()


def ensure_workbook(path):
    p = Path(path)
    if not p.exists():
        raise SystemExit(f"Workbook not found: {p}\nRun: python build_workbook.py")
    return p
