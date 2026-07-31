"""Windows recalc helper: drives LibreOffice headless to compute every formula,
writes the cached values back into the workbook, and reports any formula errors.

    python recalc_win.py dgx_llm_benchmark.xlsx
"""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"
ERRORS = ("#REF!", "#NAME?", "#VALUE!", "#DIV/0!", "#N/A", "#NULL!", "#NUM!", "Err:")


def main(path):
    src = Path(path).resolve()
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        profile = (td / "profile").as_uri()
        work = td / src.name
        shutil.copy2(src, work)
        cmd = [SOFFICE, "--headless", "--norestore", "--nolockcheck", "--nofirststartwizard",
               f"-env:UserInstallation={profile}",
               "--convert-to", "xlsx:Calc MS Excel 2007 XML",
               "--outdir", str(td / "out"), str(work)]
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
        out = td / "out" / src.name
        if not out.exists():
            print(json.dumps({"error": "conversion failed", "stdout": res.stdout, "stderr": res.stderr}))
            return 1
        shutil.copy2(out, src)

    wbf = load_workbook(src)                 # formulas
    wbv = load_workbook(src, data_only=True)  # cached values
    total_formulas = 0
    found = {}
    for name in wbf.sheetnames:
        sf, sv = wbf[name], wbv[name]
        for row in sf.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    total_formulas += 1
                    v = sv[c.coordinate].value
                    if isinstance(v, str) and any(v.startswith(e) for e in ERRORS):
                        found.setdefault(v, []).append(f"{name}!{c.coordinate}")
    total_errors = sum(len(v) for v in found.values())
    print(json.dumps({
        "status": "errors_found" if total_errors else "success",
        "total_formulas": total_formulas,
        "total_errors": total_errors,
        "error_summary": {k: v[:20] for k, v in found.items()},
    }, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else "dgx_llm_benchmark.xlsx"))
